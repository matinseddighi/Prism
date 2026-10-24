# -*- coding: utf-8 -*-
"""
Tests for convert_writer.py module
"""

import os
import sys
import tempfile
import openpyxl
import pandas as pd
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from engine.convert_writer import (
    generate_single_convert_workbook,
    generate_combined_convert_workbook,
    generate_requested_brands_workbook,
    generate_determining_workbook,
    generate_clipboard_tsv,
    export_all_converts,
    ERP_SHEET_NAME,
    ERP_HEADERS
)


def test_generate_single_convert():
    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = os.path.join(tmpdir, "هزینه ستاد.xlsx")
        rows = [
            (946, 9465, "9907052"),
            (946, 9462, "9907053")
        ]
        generate_single_convert_workbook(rows, out_path)
        assert os.path.exists(out_path)

        wb = openpyxl.load_workbook(out_path)
        assert wb.sheetnames == [ERP_SHEET_NAME]
        ws = wb.active
        assert ws.views.sheetView[0].rightToLeft is True
        rows_read = list(ws.iter_rows(values_only=True))
        assert rows_read[0] == ERP_HEADERS
        assert rows_read[1] == (946, 9465, 9907052)
        assert rows_read[2] == (946, 9462, 9907053)


def test_generate_combined_convert():
    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = os.path.join(tmpdir, "کانورت تجمیعی.xlsx")
        setad_rows = [(946, 9465, "9907052")]
        sood_rows = [(940, 9007, "9907052")]
        generate_combined_convert_workbook(setad_rows, sood_rows, out_path)
        assert os.path.exists(out_path)

        wb = openpyxl.load_workbook(out_path)
        assert wb.sheetnames == ["هزینه ستاد (946)", "سود فروشگاه (940)"]
        ws1 = wb["هزینه ستاد (946)"]
        ws2 = wb["سود فروشگاه (940)"]
        assert ws1.views.sheetView[0].rightToLeft is True
        assert ws2.views.sheetView[0].rightToLeft is True
        assert list(ws1.iter_rows(values_only=True))[1] == (946, 9465, 9907052)
        assert list(ws2.iter_rows(values_only=True))[1] == (940, 9007, 9907052)


def test_generate_clipboard_tsv():
    rows = [
        (946, 9465, "9907052"),
        (946, 9462, "9907053")
    ]
    tsv = generate_clipboard_tsv(rows)
    lines = tsv.strip().split("\n")
    assert len(lines) == 3
    assert lines[0] == "کد دسته اصلي (*)\tکد دسته فرعي (*)\tکد کالا (*)"
    assert lines[1] == "946\t9465\t9907052"
    assert lines[2] == "946\t9462\t9907053"


def test_export_all_converts():
    with tempfile.TemporaryDirectory() as tmp_export, tempfile.TemporaryDirectory() as tmp_downloads:
        missing_items = [
            {
                "code": "9907052",
                "name": "پنیر پیتزا دالیا",
                "brand": "دالیا",
                "category": "لبنیات",
                "determined_setad": 5,
                "determined_sood": 7,
                "calc_price": 50000,
                "cons_price_no_tax": 60000,
                "supplier_discount": 10,
                "proposed_convert_discount": 7,
                "status": "سالم",
                "price_gap": 5000,
                "gross_margin_total": 10000,
                "store_share_pct": 70
            }
        ]
        new_goods = [
            {
                "code": "9901111",
                "name": "کالای جدید",
                "brand": "برند",
                "category": "گروه",
                "opening_date": "1405/07/30",
                "calc_price": 10000
            }
        ]

        result = export_all_converts(
            date_str="1405.07.30",
            source_filename="تست.xlsx",
            missing_items=missing_items,
            requested_data=None,
            new_goods=new_goods,
            total_goods_count=100,
            sood_setad_data=None,
            export_dir=tmp_export,
            downloads_dir=tmp_downloads
        )

        assert result["success"] is True
        assert result["setad_count"] == 1
        assert result["sood_count"] == 1
        assert os.path.exists(os.path.join(tmp_export, "هزینه ستاد.xlsx"))
        assert os.path.exists(os.path.join(tmp_export, "سود فروشگاه کانورت.xlsx"))
        assert os.path.exists(os.path.join(tmp_export, "کانورت هزینه ستاد و سود فروشگاه 1405.07.30.xlsx"))
        assert os.path.exists(os.path.join(tmp_export, "تعیین سود و هزینه کالاهای فاقد مقدار 1405.07.30.xlsx"))
        # Check files were copied to downloads
        assert os.path.exists(os.path.join(tmp_downloads, "هزینه ستاد.xlsx"))
