# -*- coding: utf-8 -*-
"""
Tests for ShelfMaster Convert Module (engine/convert.py)
Validates TSV clipboard generation, ERP-compliant 13-column Excel output,
SOP enforcement, styling, and the Anti-Lock fallback mechanism.
"""

import os
import sys
import tempfile
from unittest.mock import patch
import openpyxl
import pytest

# Ensure ShelfMaster root is on sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from engine.convert import (
    HEADERS,
    generate_tsv_clipboard,
    generate_convert_file,
)
from engine.rules import STORE_WAREHOUSE_MAP, PERMANENT_ZERO_CODES


def test_headers_exact_13_columns_and_arabic_yeh():
    """Verify exact 13 columns and exact 5 ERP headers with Arabic Yeh (\u064a)."""
    assert len(HEADERS) == 13
    assert HEADERS[0] == "کد انبار(*)"
    assert HEADERS[1] == "کد کالا(*)"
    assert HEADERS[2] == "کف موجودي(*)"
    assert HEADERS[3] == "سقف موجودي(*)"
    assert HEADERS[4] == "نقطه ثبت سفارش(*)"
    assert HEADERS[5] == "نام محصول"
    assert HEADERS[6] == "کف قبلی"
    assert HEADERS[7] == "سقف قبلی"
    assert HEADERS[8] == "تغییرات سقف / وضعیت تعادل"
    assert HEADERS[9] == "تعداد بچ مادر"
    assert HEADERS[10] == "وضعیت کالا"
    assert HEADERS[11] == "موجودی قفسه شعبه"
    assert HEADERS[12] == "موجودی قابل ارسال"

    # Strict check: character at index 8 of "کف موجودي(*)" must be Arabic Yeh '\u064a'
    assert "\u064a" in HEADERS[2], "Header 3 must contain Arabic Yeh (ي)"
    assert "\u064a" in HEADERS[3], "Header 4 must contain Arabic Yeh (ي)"


def test_generate_tsv_clipboard_basic():
    """Unit test for generate_tsv_clipboard (columns, delimiter, row count)."""
    items = [
        {"code": 9907092, "min": 25, "max": 100},
        {"code": 9901234, "min": 10, "max": 30},
    ]
    tsv = generate_tsv_clipboard(items, store_code=3401)
    lines = tsv.strip().split("\n")
    assert len(lines) == 2

    for line in lines:
        cols = line.split("\t")
        assert len(cols) == 5, f"Each row in TSV must have 5 tab-separated columns, got {cols}"
        assert cols[0] == "3401"
        assert cols[4] == "0"

    row1 = lines[0].split("\t")
    assert row1[1] == "9907092"
    assert row1[2] == "25"
    assert row1[3] == "100"

    row2 = lines[1].split("\t")
    assert row2[1] == "9901234"
    assert row2[2] == "10"
    assert row2[3] == "30"


def test_generate_tsv_clipboard_sop_enforcement():
    """Verify TSV generator enforces Max > Min and Zero Mandated rules."""
    items = [
        # Max <= Min -> Auto-bumped Max
        {"code": 9905555, "min": 10, "max": 5, "batch": 24},
        # Permanent zero detergent (10L Golrang)
        {"code": 9900910, "min": 10, "max": 20, "name": "مایع دستشویی ۱۰ لیتری گلرنگ"},
    ]
    tsv = generate_tsv_clipboard(items, store_code=3601)
    lines = tsv.strip().split("\n")
    assert len(lines) == 2

    cols1 = lines[0].split("\t")
    assert int(cols1[2]) == 10
    assert int(cols1[3]) > 10, "Max must be greater than Min in TSV output"

    cols2 = lines[1].split("\t")
    assert int(cols2[2]) == 0, "Zero mandated item must have Min=0 in TSV"
    assert int(cols2[3]) == 0, "Zero mandated item must have Max=0 in TSV"


def test_generate_convert_file_structure_and_readback():
    """Test generating full Excel convert file and reading back with openpyxl."""
    with tempfile.TemporaryDirectory() as tmpdir:
        items = [
            {
                "code": 9907092,
                "name": "روغن آفتابگردان ۸۱۰ گرمی لادن",
                "min": 24,
                "max": 72,
                "old_min": 12,
                "old_max": 36,
                "batch": 12,
                "status": "فعال",
                "shelf_qty": 10,
                "free_qty": 500,
            },
            {
                # Needs Max bump: min=15, max=10, batch=20
                "code": 9901111,
                "name": "تن ماهی ۱۸۰ گرمی تحفه",
                "min": 15,
                "max": 10,
                "old_min": 15,
                "old_max": 10,
                "batch": 20,
                "status": "فعال",
                "shelf_qty": 5,
                "free_qty": 200,
            },
            {
                # Permanent zero item
                "code": 9902937,
                "name": "سبد سفید ۲۵ کیلویی پالود پلاست",
                "min": 5,
                "max": 10,
                "batch": 1,
                "status": "فعال",
                "shelf_qty": 2,
                "free_qty": 0,
            },
        ]

        res = generate_convert_file(
            store_name="فضل",
            items=items,
            target_date="1405.07.29",
            output_dirs=[tmpdir],
        )

        assert res["count"] == 3
        assert res["store_code"] == 3401
        assert "paths" in res and len(res["paths"]) == 1
        saved_file = res["paths"][0]
        assert os.path.isfile(saved_file)

        # Read back with openpyxl
        wb = openpyxl.load_workbook(saved_file)
        ws = wb.active
        assert ws.title == "کانورت فضل"
        assert ws.views.sheetView[0].rightToLeft is True

        # Check Row 1 headers
        read_headers = [ws.cell(row=1, column=c).value for c in range(1, 14)]
        assert read_headers == HEADERS

        # Check Styling on Row 1
        assert ws.row_dimensions[1].height == 28
        for c in range(1, 6):
            cell = ws.cell(row=1, column=c)
            # Navy Blue #1B365D
            assert cell.fill.start_color.rgb in ("001B365D", "FF1B365D", "1B365D")
            assert cell.font.bold is True
            assert cell.font.color.rgb in ("00FFFFFF", "FFFFFFFF", "FFFFFF")
        for c in range(6, 14):
            cell = ws.cell(row=1, column=c)
            # Steel Blue #335C8D
            assert cell.fill.start_color.rgb in ("00335C8D", "FF335C8D", "335C8D")

        # Row 2: Standard Item
        r2_wh = ws.cell(row=2, column=1).value
        r2_code = ws.cell(row=2, column=2).value
        r2_min = ws.cell(row=2, column=3).value
        r2_max = ws.cell(row=2, column=4).value
        r2_order = ws.cell(row=2, column=5).value
        assert r2_wh == 3401
        assert r2_code == 9907092
        assert r2_min == 24
        assert r2_max == 72
        assert r2_order == 0

        # Row 3: Max > Min Bumped Item
        r3_min = ws.cell(row=3, column=3).value
        r3_max = ws.cell(row=3, column=4).value
        assert r3_max > r3_min
        assert r3_min == 15
        assert r3_max == 20  # carton 20 > 15
        # Check soft green highlight on updated max (#E2EFDA)
        cell_r3_max = ws.cell(row=3, column=4)
        assert cell_r3_max.fill.start_color.rgb in ("00E2EFDA", "FFE2EFDA", "E2EFDA")

        # Row 4: Permanent zero item (Palood basket)
        r4_min = ws.cell(row=4, column=3).value
        r4_max = ws.cell(row=4, column=4).value
        assert r4_min == 0
        assert r4_max == 0
        r4_note = ws.cell(row=4, column=9).value
        assert "صفر" in str(r4_note)

        # Verify alignment
        for r in range(2, 5):
            assert ws.cell(row=r, column=6).alignment.horizontal == "right"
            assert ws.cell(row=r, column=9).alignment.horizontal == "right"
            assert ws.cell(row=r, column=1).alignment.horizontal == "center"
            assert ws.cell(row=r, column=3).alignment.horizontal == "center"


def test_anti_lock_fallback_mechanism():
    """Verify that PermissionError triggers automatic fallback to '(اصلاحی).xlsx'."""
    with tempfile.TemporaryDirectory() as tmpdir:
        items = [{"code": 9901234, "min": 10, "max": 20, "name": "کالای تستی"}]
        target_fname = "کانورت کف و سقف بوژان (1405.07.29).xlsx"
        locked_path = os.path.join(tmpdir, target_fname)

        original_save = openpyxl.Workbook.save

        def fake_save(self, filename):
            # If attempting to save the exact locked_path, simulate Excel locking it
            if os.path.normpath(str(filename)) == os.path.normpath(locked_path):
                raise PermissionError(f"[WinError 32] The process cannot access the file: '{filename}'")
            return original_save(self, filename)

        with patch.object(openpyxl.Workbook, "save", side_effect=fake_save, autospec=True):
            res = generate_convert_file(
                store_name="بوژان",
                items=items,
                target_date="1405.07.29",
                output_dirs=[tmpdir],
            )

        assert len(res["paths"]) == 1
        saved_file = res["paths"][0]
        assert "(اصلاحی).xlsx" in saved_file
        assert os.path.isfile(saved_file)
        assert res["store_code"] == 4101


def test_double_anti_lock_fallback():
    """Verify that if both original and '(اصلاحی).xlsx' are locked, it falls back to '(اصلاحی 2).xlsx'."""
    with tempfile.TemporaryDirectory() as tmpdir:
        items = [{"code": 9901234, "min": 10, "max": 20}]
        target_fname = "کانورت کف و سقف الهیه (1405.07.29).xlsx"
        locked_path1 = os.path.join(tmpdir, target_fname)
        locked_path2 = os.path.join(tmpdir, "کانورت کف و سقف الهیه (1405.07.29) (اصلاحی).xlsx")

        original_save = openpyxl.Workbook.save

        def fake_save(self, filename):
            fn_norm = os.path.normpath(str(filename))
            if fn_norm in (os.path.normpath(locked_path1), os.path.normpath(locked_path2)):
                raise PermissionError(f"[WinError 32] Locked: '{filename}'")
            return original_save(self, filename)

        with patch.object(openpyxl.Workbook, "save", side_effect=fake_save, autospec=True):
            res = generate_convert_file(
                store_name="الهیه",
                items=items,
                target_date="1405.07.29",
                output_dirs=[tmpdir],
            )

        assert len(res["paths"]) == 1
        saved_file = res["paths"][0]
        assert "(اصلاحی 2).xlsx" in saved_file
        assert os.path.isfile(saved_file)
        assert res["store_code"] == 4201


def test_persian_alias_store_names_and_code_resolution():
    """Verify store name aliases with Persian digits and arabic yeh resolve to correct 4-digit codes."""
    with tempfile.TemporaryDirectory() as tmpdir:
        test_stores = [
            ("۱۷ شهریور", 4301),
            ("17 شهریور", 4301),
            ("فردوسی", 3601),
            ("الینا", 2901),
            ("الينا", 2901),
            ("بهشتی", 3301),
            ("مدرس", 3501),
            ("رحمت آباد", 3901),
        ]
        for sname, expected_code in test_stores:
            res = generate_convert_file(
                store_name=sname,
                items=[{"code": 100, "min": 5, "max": 10}],
                output_dirs=[tmpdir],
            )
            assert res["store_code"] == expected_code, f"Store '{sname}' expected {expected_code}, got {res['store_code']}"


def test_slow_moving_item_cap_in_convert():
    """Verify Section 6 rule: slow moving item with carton <= 12 capped at 1 carton."""
    with tempfile.TemporaryDirectory() as tmpdir:
        items = [
            {
                "code": 9908888,
                "name": "خمیر دندان مریدنت ۱۰۰ میل",
                "min": 4,
                "max": 36,  # User or raw was 36 (3 cartons of 12)
                "batch": 12,
                "is_slow_moving": True,
            }
        ]
        res = generate_convert_file(
            store_name="فضل",
            items=items,
            output_dirs=[tmpdir],
        )
        saved_file = res["paths"][0]
        wb = openpyxl.load_workbook(saved_file)
        ws = wb.active

        adj_max = ws.cell(row=2, column=4).value
        assert adj_max == 12, f"Slow moving item with carton=12 must be capped at 12, got {adj_max}"
        note = ws.cell(row=2, column=9).value
        assert "کندگرد" in str(note)


def test_empty_items_list():
    """Verify empty items list produces valid workbook with headers and empty TSV."""
    with tempfile.TemporaryDirectory() as tmpdir:
        res = generate_convert_file(
            store_name="پژواک",
            items=[],
            output_dirs=[tmpdir],
        )
        assert res["count"] == 0
        assert res["tsv"] == ""
        assert res["store_code"] == 3001  # پژواک code
        wb = openpyxl.load_workbook(res["paths"][0])
        ws = wb.active
        assert ws.max_row == 1
        assert [ws.cell(row=1, column=c).value for c in range(1, 14)] == HEADERS


def test_persian_key_fallbacks_in_items():
    """Verify items dictionary supports Persian header keys seamlessly."""
    with tempfile.TemporaryDirectory() as tmpdir:
        items = [
            {
                "کد کالا": 9901234,
                "نام کالا": "روغن حیات ۱.۵ لیتری",
                "کف": 10,
                "سقف": 30,
                "کف قبلی": 5,
                "سقف قبلی": 15,
                "تعداد بچ مادر": 6,
                "وضعیت کالا": "فعال",
                "موجودی قفسه شعبه": 8,
                "موجودی قابل ارسال": 120,
            }
        ]
        res = generate_convert_file(
            store_name="فضل",
            items=items,
            output_dirs=[tmpdir],
        )
        wb = openpyxl.load_workbook(res["paths"][0])
        ws = wb.active
        assert ws.cell(row=2, column=1).value == 3401
        assert ws.cell(row=2, column=2).value == 9901234
        assert ws.cell(row=2, column=3).value == 10
        assert ws.cell(row=2, column=4).value == 30
        assert ws.cell(row=2, column=6).value == "روغن حیات ۱.۵ لیتری"
        assert ws.cell(row=2, column=7).value == 5
        assert ws.cell(row=2, column=8).value == 15
        assert ws.cell(row=2, column=10).value == 6


def test_previous_columns_use_prev_values_and_highlight_changes():
    """Edited table rows / presets / audit-fix send prev_min/prev_max: Excel «کف قبلی/سقف قبلی» must show them,
    and only really-changed cells are green."""
    with tempfile.TemporaryDirectory() as tmpdir:
        items = [
            {"wh_name": "فضل", "code": 9900001, "name": "کالای ویرایش‌شده", "min": 4, "max": 12,
             "prev_min": 4, "prev_max": 12, "new_min": 6, "new_max": 12, "batch": 6},
        ]
        res = generate_convert_file(store_name="فضل", items=items, output_dirs=[tmpdir])
        ws = openpyxl.load_workbook(res["paths"][0]).active
        assert ws.cell(row=2, column=3).value == 6    # new floor
        assert ws.cell(row=2, column=4).value == 12   # ceiling unchanged
        assert ws.cell(row=2, column=7).value == 4    # previous floor (was wrongly 6 before)
        assert ws.cell(row=2, column=8).value == 12
        assert ws.cell(row=2, column=3).fill.start_color.rgb.endswith("E2EFDA")
        assert not ws.cell(row=2, column=4).fill.start_color.rgb.endswith("E2EFDA")


def test_row_store_code_zero_uses_store_or_name():
    """An explicit store_code of 0 (e.g. AI item without store) must not end up in the ERP file."""
    tsv = generate_tsv_clipboard([{"code": 1, "min": 1, "max": 3, "store_code": 0}], 3501)
    assert tsv.startswith("3501\t")
    tsv = generate_tsv_clipboard([{"code": 1, "min": 1, "max": 3, "store_code": 0, "wh_name": "بهشتی"}], 0)
    assert tsv.startswith("3301\t")


def _xlsx_bytes(rows):
    import io
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_custom_excel_reads_panel_own_convert_format():
    """Re-loading a convert file: «موجودی قفسه شعبه» must not be taken as the store; the store comes from «کد انبار(*)»."""
    from engine.custom_excel import inspect_and_convert_custom_excel
    data = _xlsx_bytes([HEADERS, [3501, 9900001, 6, 12, 0, "کالای تست", 4, 12, "", 6, "فعال", 3, 0]])
    res = inspect_and_convert_custom_excel(data)
    assert res["ok"] is True
    assert res["store_code"] == 3501
    assert res["store_name"] == "مدرس"
    assert res["items"][0]["store_code"] == 3501
    assert res["items"][0]["min"] == 6


def test_custom_excel_multi_store_rows_keep_their_own_code():
    from engine.custom_excel import inspect_and_convert_custom_excel
    data = _xlsx_bytes([
        ["نام فروشگاه", "کد کالا", "نام کالا", "کف فعلی قفسه", "سقف فعلی قفسه", "موجودی قفسه شعبه"],
        ["فضل", 9900001, "الف", 2, 6, 1],
        ["بهشتی", 9900002, "ب", 3, 9, 0],
    ])
    res = inspect_and_convert_custom_excel(data)
    assert [it["store_code"] for it in res["items"]] == [3401, 3301]
    assert res["tsv"].splitlines()[1].startswith("3301\t9900002")


def test_custom_excel_unknown_store_is_refused():
    from engine.custom_excel import inspect_and_convert_custom_excel
    data = _xlsx_bytes([["کد کالا", "کف", "سقف"], [9900001, 2, 6]])
    res = inspect_and_convert_custom_excel(data)
    assert res["ok"] is False
    assert "کد انبار" in res["error"]


def test_default_output_dirs_redirected_in_tests():
    """conftest sets SHELFMASTER_OUTPUT_DIR, so default exports never touch D:\\ or Downloads during tests."""
    from engine.convert import default_output_dirs
    assert default_output_dirs("1405.08.01") == [os.environ["SHELFMASTER_OUTPUT_DIR"]]
