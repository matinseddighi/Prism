# -*- coding: utf-8 -*-
"""
ShelfMaster Engine - Convert & TSV Clipboard Generation Module
Produces ERP-standard 13-column Excel files and 5-column clipboard text
with strict SOP rule enforcement, dual/triple path sync, and Anti-Lock protection.
"""

from __future__ import annotations

import os
import re
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from .rules import (
        STORE_WAREHOUSE_MAP,
        get_store_warehouse_code,
        is_zero_mandated,
        validate_and_adjust_limits,
    )
except (ImportError, ValueError):
    from rules import (
        STORE_WAREHOUSE_MAP,
        get_store_warehouse_code,
        is_zero_mandated,
        validate_and_adjust_limits,
    )

# 1. 13 Standard Integrated Columns per SOP (Arabic Yeh \u064a in موجودي(*))
HEADERS: list[str] = [
    "کد انبار(*)",
    "کد کالا(*)",
    "کف موجودي(*)",
    "سقف موجودي(*)",
    "نقطه ثبت سفارش(*)",
    "نام محصول",
    "کف قبلی",
    "سقف قبلی",
    "تغییرات سقف / وضعیت تعادل",
    "تعداد بچ مادر",
    "وضعیت کالا",
    "موجودی قفسه شعبه",
    "موجودی قابل ارسال",
]


def _safe_int(val: Any, default: int = 0) -> int:
    """Convert value safely to int."""
    if val is None or val == "":
        return default
    try:
        return int(round(float(str(val).strip())))
    except (ValueError, TypeError):
        return default


def _safe_float(val: Any, default: float = 0.0) -> float:
    """Convert value safely to float."""
    if val is None or val == "":
        return default
    try:
        return float(str(val).strip())
    except (ValueError, TypeError):
        return default


def _get_item_val(it: dict[str, Any], keys: list[str], default: Any = None) -> Any:
    """Retrieve first matching non-null value from item dict across alias keys."""
    for k in keys:
        if k in it and it[k] is not None:
            return it[k]
    return default


STORE_CODE_KEYS = ["store_code", "wh_code", "کد انبار(*)", "کد انبار"]
# مقدار قبلی (ستون «کف قبلی/سقف قبلی» اکسل): old_* از اکسل سفارشی/دستیار، prev_* از سناریو، اصلاح خودکار و جدول
OLD_MIN_KEYS = ["old_min", "prev_min", "کف قبلی", "min_num", "کف فعلی قفسه", "کف فعلی"]
OLD_MAX_KEYS = ["old_max", "prev_max", "سقف قبلی", "max_num", "سقف فعلی قفسه", "سقف فعلی"]


def _row_store_code(it: dict[str, Any], store_code: int) -> int:
    """Row's own warehouse code if it has a valid one (> 0), else its store name's code, else the export store code."""
    for k in STORE_CODE_KEYS:
        c = _safe_int(it.get(k), 0)
        if c > 0:
            return c
    for k in ("wh_name", "store"):
        if it.get(k):
            c = get_store_warehouse_code(str(it[k]))
            if c > 0:
                return c
    return store_code


def default_output_dirs(target_date: str) -> list[str]:
    """Daily SOP folders. SHELFMASTER_OUTPUT_DIR (used by tests) redirects every write to one folder."""
    override = os.environ.get("SHELFMASTER_OUTPUT_DIR")
    if override:
        return [override]
    return [
        os.path.join(r"D:\1.sed daily 405\3. کانورت کف و سقف", target_date),
        os.path.join(r"D:\personal\projects\high and low", target_date),
        os.path.join(os.path.expanduser("~"), "Downloads"),
    ]


def generate_tsv_clipboard(items: list[dict[str, Any]], store_code: int) -> str:
    """
    Generate tab-separated 5 columns for direct clipboard pasting into ERP software.
    Columns: [store_code, code, min, max, 0].
    Enforces SOP rules (Max > Min invariant and zero-mandated items).
    """
    lines: list[str] = []
    for it in items:
        code = _safe_int(_get_item_val(it, ["code", "کد کالا", "کد", "کد کالا(*)"], 0))
        name = str(_get_item_val(it, ["name", "نام محصول", "نام کالا", "شرح کالا", "عنوان کالا"], ""))
        carton = _safe_float(
            _get_item_val(it, ["batch", "carton", "تعداد بچ مادر", "بچ مادر", "بچ"], 1.0),
            default=1.0,
        )
        if carton <= 0:
            carton = 1.0

        raw_min = _get_item_val(it, ["new_min", "min", "کف", "کف موجودي(*)", "کف موجودی(*)", "کف جدید"], 0)
        raw_max = _get_item_val(it, ["new_max", "max", "سقف", "سقف موجودي(*)", "سقف موجودی(*)", "سقف جدید"], 0)
        is_slow = bool(it.get("is_slow_moving", False))

        if is_zero_mandated(code, name):
            adj_min, adj_max = 0, 0
        else:
            adj_min, adj_max, _ = validate_and_adjust_limits(
                raw_min, raw_max, carton, is_slow_moving=is_slow
            )

        row_store_code = _row_store_code(it, store_code)
        lines.append(f"{row_store_code}\t{code}\t{adj_min}\t{adj_max}\t0")

    return "\n".join(lines)


def generate_convert_file(
    store_name: str,
    items: list[dict[str, Any]],
    target_date: str = "1405.07.29",
    output_dirs: list[str] | None = None,
    file_label: str | None = None,
) -> dict[str, Any]:
    """
    Generate standard ERP 13-column single-sheet convert workbook with RTL view,
    dual/triple path synchronization, SOP enforcement, and Anti-Lock fallback protection.

    Parameters:
        store_name: Store name (e.g. 'فضل', 'بوژان', 'فردوسی شمالی')
        items: List of product dictionaries
        target_date: Shamsi date string for directory and filename
        output_dirs: Optional list of directories to save to (defaults to daily SOP paths)

    Returns:
        dict: {"filename": base_name, "paths": saved_paths, "count": len(items), "tsv": tsv_text, "store_code": store_code}
    """
    # 1. Resolve store warehouse 4-digit code
    store_code = get_store_warehouse_code(store_name)
    if store_code == 0:
        # Check if store_name is already a numeric code
        try:
            store_code = int(str(store_name).strip())
        except (ValueError, TypeError):
            # Check if any item contains store_code
            for it in items:
                sc = _safe_int(_get_item_val(it, ["store_code", "wh_code", "کد انبار(*)", "کد انبار"]))
                if sc > 0:
                    store_code = sc
                    break

    # 2. Setup OpenPyXL Workbook
    wb = openpyxl.Workbook()
    ws = wb.active

    # Clean title for Excel sheet limit (max 31 characters, no invalid symbols)
    clean_title = re.sub(r"[:\\/?*\[\]]", "", str(store_name)).strip()[:20]
    ws.title = f"کانورت {clean_title}"
    ws.views.sheetView[0].rightToLeft = True

    # 3. Define Styles
    font_hdr = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Tahoma", size=9)
    font_bold = Font(name="Tahoma", size=9, bold=True)

    fill_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_steel = PatternFill(start_color="335C8D", end_color="335C8D", fill_type="solid")
    fill_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # 4. Write Header Row (Row 1)
    ws.row_dimensions[1].height = 28
    ws.append(HEADERS)
    for col_idx in range(1, 14):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_hdr
        cell.fill = fill_navy if col_idx <= 5 else fill_steel
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # 5. Populate Data Rows
    processed_items: list[dict[str, Any]] = []

    for r_idx, it in enumerate(items, start=2):
        ws.row_dimensions[r_idx].height = 22

        code = _safe_int(_get_item_val(it, ["code", "کد کالا", "کد", "کد کالا(*)"], 0))
        name = str(_get_item_val(it, ["name", "نام محصول", "نام کالا", "شرح کالا", "عنوان کالا"], ""))
        carton = _safe_float(
            _get_item_val(it, ["batch", "carton", "تعداد بچ مادر", "بچ مادر", "بچ"], 1.0),
            default=1.0,
        )
        if carton <= 0:
            carton = 1.0

        raw_min = _get_item_val(it, ["new_min", "min", "کف", "کف موجودي(*)", "کف موجودی(*)", "کف جدید"], 0)
        raw_max = _get_item_val(it, ["new_max", "max", "سقف", "سقف موجودي(*)", "سقف موجودی(*)", "سقف جدید"], 0)
        is_slow = bool(it.get("is_slow_moving", False))

        # Enforce SOP rules
        if is_zero_mandated(code, name):
            adj_min, adj_max = 0, 0
            sop_note = "الزام سقف و کف صفر طبق SOP (اقلام صنعتی/شوینده ۱۰ لیتری/سبد)"
        else:
            adj_min, adj_max, sop_note = validate_and_adjust_limits(
                raw_min, raw_max, carton, is_slow_moving=is_slow
            )

        # Resolve previous values for comparison and highlight
        old_min_raw = _get_item_val(it, OLD_MIN_KEYS)
        old_max_raw = _get_item_val(it, OLD_MAX_KEYS)
        # «نامشخص» (old value unknown in a custom Excel) counts as no previous value
        if old_min_raw is not None and str(old_min_raw).strip() in ("", "نامشخص"):
            old_min_raw = None
        if old_max_raw is not None and str(old_max_raw).strip() in ("", "نامشخص"):
            old_max_raw = None

        has_explicit_old_min = old_min_raw is not None
        has_explicit_old_max = old_max_raw is not None

        old_min = _safe_int(old_min_raw, default=adj_min)
        old_max = _safe_int(old_max_raw, default=adj_max)

        # Determine update highlights
        if has_explicit_old_min:
            min_updated = adj_min != old_min
        else:
            min_updated = adj_min != _safe_int(raw_min)

        if has_explicit_old_max:
            max_updated = adj_max != old_max
        else:
            max_updated = adj_max != _safe_int(raw_max)

        change_note = (
            _get_item_val(it, ["change_note", "diff_note", "تغییرات سقف / وضعیت تعادل", "تغییرات سقف", "توضیحات"])
            or sop_note
        )

        batch_int = int(round(carton)) if carton == int(round(carton)) else round(carton, 2)
        status_val = str(_get_item_val(it, ["status", "وضعیت کالا", "وضعیت کالا در فروشگاه", "وضعیت"], "فعال"))
        shelf_qty = _safe_int(_get_item_val(it, ["shelf_qty", "موجودی قفسه شعبه", "موجودی قفسه", "موجودی انبار قفسه"], 0))
        free_qty = _safe_int(_get_item_val(it, ["free_qty", "موجودی قابل ارسال", "موجودی آزاد"], 0))
        row_store_code = _row_store_code(it, store_code)

        row_vals = [
            row_store_code,
            code,
            adj_min,
            adj_max,
            0,
            name,
            old_min,
            old_max,
            change_note,
            batch_int,
            status_val,
            shelf_qty,
            free_qty,
        ]

        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = font_bold if col_idx in (3, 4) else font_data
            cell.border = thin_border

            # Alignment: Col 6 & 9 right-aligned, others centered
            if col_idx in (6, 9):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Soft green highlight on updated Min / Max
            if col_idx == 3 and min_updated:
                cell.fill = fill_green
            elif col_idx == 4 and max_updated:
                cell.fill = fill_green
            elif r_idx % 2 == 1:
                cell.fill = fill_zebra

        processed_items.append({
            "store_code": row_store_code,
            "code": code,
            "min": adj_min,
            "max": adj_max,
            "name": name,
        })

    # 6. Set Column Widths
    col_widths: dict[int, int] = {
        1: 14,
        2: 15,
        3: 16,
        4: 16,
        5: 18,
        6: 42,
        7: 12,
        8: 12,
        9: 50,
        10: 16,
        11: 14,
        12: 18,
        13: 18,
    }
    for c_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = width

    # 7. File Saving & Anti-Lock Fallback
    # file_label: e.g. «فضل - ویرایش» so a partial export doesn't overwrite the day's full store file
    label = re.sub(r'[\\/:*?"<>|]', "", str(file_label or store_name)).strip() or "بدون نام"
    base_fname = f"کانورت کف و سقف {label} ({target_date}).xlsx"

    dirs = output_dirs if output_dirs is not None else default_output_dirs(target_date)
    target_paths = [os.path.join(d, base_fname) for d in dirs]

    saved_paths: list[str] = []
    for p in target_paths:
        try:
            parent_dir = os.path.dirname(p)
            if parent_dir:
                os.makedirs(parent_dir, exist_ok=True)

            try:
                wb.save(p)
                saved_paths.append(p)
            except PermissionError:
                # Anti-Lock mechanism: file is open in Excel, fallback to (اصلاحی)
                alt = p.replace(".xlsx", " (اصلاحی).xlsx")
                try:
                    wb.save(alt)
                    saved_paths.append(alt)
                except PermissionError:
                    alt2 = p.replace(".xlsx", " (اصلاحی 2).xlsx")
                    wb.save(alt2)
                    saved_paths.append(alt2)
        except (OSError, IOError):
            # Gracefully ignore unreachable paths (e.g. non-existent drive letter during tests)
            pass

    # Safety guarantee: ensure at least one path is saved if all target_paths were unwriteable
    if not saved_paths:
        fallback_path = os.path.abspath(os.path.join(dirs[0] if dirs and os.path.isdir(dirs[0]) else ".", base_fname))
        try:
            wb.save(fallback_path)
            saved_paths.append(fallback_path)
        except PermissionError:
            alt = fallback_path.replace(".xlsx", " (اصلاحی).xlsx")
            wb.save(alt)
            saved_paths.append(alt)

    tsv_content = generate_tsv_clipboard(processed_items, store_code)

    return {
        "filename": os.path.basename(saved_paths[0]),
        "paths": saved_paths,
        "count": len(items),
        "tsv": tsv_content,
        "store_code": store_code,
    }
